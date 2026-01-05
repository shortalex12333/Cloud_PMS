"""
CelesteOS Document Extraction Service
Extracts text from any document type - PDF, DOCX, XLSX, CSV, TXT, etc.
"""

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import pdfplumber
import httpx
import io
import os
import csv
import json

app = FastAPI(
    title="CelesteOS Extraction Service",
    description="Universal document text extraction",
    version="2.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- Models ---

class ExtractRequest(BaseModel):
    storage_path: str
    content_type: str
    filename: str
    yacht_id: str
    document_id: str
    service_key: str = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InZ6c29oYXZ0dW90b2NncmZrZnlkIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIsImlhdCI6MTc2MzU5Mjg3NSwiZXhwIjoyMDc5MTY4ODc1fQ.fC7eC_4xGnCHIebPzfaJ18pFMPKgImE7BuN0I3A-pSY"
    # Pass-through metadata
    doc_type: str = None
    system_tag: str = None

class ExtractionResponse(BaseModel):
    text: str
    pages: int = 0
    chars: int = 0
    status: str = "success"
    filename: str = ""
    content_type: str = ""
    yacht_id: str = ""
    document_id: str = ""
    doc_type: str = None
    system_tag: str = None

# --- Health ---

@app.get("/")
def health():
    return {"status": "ok", "service": "extraction", "version": "2.0.0"}

@app.get("/health")
def health_check():
    return {"status": "healthy", "version": "2.0.0", "supported_types": [
        "application/pdf",
        "text/plain", "text/csv", "text/html", "text/xml",
        "application/json", "application/xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ]}

# --- Main Extraction Endpoint ---

SUPABASE_STORAGE_BASE = "https://vzsohavtuotocgrfkfyd.supabase.co/storage/v1/object"

@app.post("/extract", response_model=ExtractionResponse)
async def extract(request: ExtractRequest):
    """
    Main extraction endpoint.
    Fetches file from Supabase storage and extracts text based on content_type.
    """
    try:
        # Build storage URL
        storage_url = f"{SUPABASE_STORAGE_BASE}/{request.storage_path}"
        
        # Fetch file from Supabase
        headers = {
            "apikey": request.service_key,
            "Authorization": f"Bearer {request.service_key}"
        }
        
        async with httpx.AsyncClient(timeout=120.0) as client:
            resp = await client.get(storage_url, headers=headers)
            if resp.status_code == 404:
                raise HTTPException(status_code=404, detail=f"File not found: {request.storage_path}")
            resp.raise_for_status()
            file_bytes = resp.content
        
        # Extract based on content type
        text, pages = await _extract_by_type(file_bytes, request.content_type, request.filename)
        
        return ExtractionResponse(
            text=text,
            pages=pages,
            chars=len(text),
            status="success",
            filename=request.filename,
            content_type=request.content_type,
            yacht_id=request.yacht_id,
            document_id=request.document_id,
            doc_type=request.doc_type,
            system_tag=request.system_tag
        )
    
    except httpx.HTTPError as e:
        raise HTTPException(status_code=502, detail=f"Failed to fetch file: {str(e)}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Extraction failed: {str(e)}")


async def _extract_by_type(file_bytes: bytes, content_type: str, filename: str) -> tuple[str, int]:
    """Route extraction based on content type."""
    
    # PDF
    if content_type == "application/pdf" or filename.lower().endswith(".pdf"):
        return _extract_pdf(file_bytes)
    
    # Excel XLSX
    if content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or filename.lower().endswith(".xlsx"):
        return _extract_xlsx(file_bytes)
    
    # Word DOCX
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or filename.lower().endswith(".docx"):
        return _extract_docx(file_bytes)
    
    # CSV
    if content_type == "text/csv" or filename.lower().endswith(".csv"):
        return _extract_csv(file_bytes)
    
    # JSON
    if content_type == "application/json" or filename.lower().endswith(".json"):
        return _extract_json(file_bytes)
    
    # XML
    if content_type in ["application/xml", "text/xml"] or filename.lower().endswith(".xml"):
        return _extract_text(file_bytes)  # Just read as text
    
    # HTML
    if content_type == "text/html" or filename.lower().endswith(".html"):
        return _extract_html(file_bytes)
    
    # Default: plain text
    return _extract_text(file_bytes)


def _extract_pdf(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from PDF using pdfplumber."""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        pages_text = []
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages_text.append(text)
        return "\n\n".join(pages_text), len(pdf.pages)


def _extract_xlsx(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from Excel XLSX."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        all_text = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                if row_text.strip():
                    rows.append(row_text)
            if rows:
                all_text.append(f"=== {sheet.title} ===\n" + "\n".join(rows))
        return "\n\n".join(all_text), len(wb.worksheets)
    except ImportError:
        # Fallback if openpyxl not available
        return "[XLSX extraction requires openpyxl]", 0


def _extract_docx(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from Word DOCX."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs), len(paragraphs)
    except ImportError:
        # Fallback if python-docx not available
        return "[DOCX extraction requires python-docx]", 0


def _extract_csv(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from CSV."""
    text = file_bytes.decode('utf-8', errors='ignore')
    reader = csv.reader(io.StringIO(text))
    rows = [" | ".join(row) for row in reader]
    return "\n".join(rows), len(rows)


def _extract_json(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from JSON - pretty print it."""
    text = file_bytes.decode('utf-8', errors='ignore')
    try:
        data = json.loads(text)
        return json.dumps(data, indent=2), 1
    except:
        return text, 1


def _extract_html(file_bytes: bytes) -> tuple[str, int]:
    """Extract text from HTML - strip tags."""
    import re
    text = file_bytes.decode('utf-8', errors='ignore')
    # Simple tag stripping
    clean = re.sub(r'<[^>]+>', ' ', text)
    clean = re.sub(r'\s+', ' ', clean).strip()
    return clean, 1


def _extract_text(file_bytes: bytes) -> tuple[str, int]:
    """Extract plain text."""
    return file_bytes.decode('utf-8', errors='ignore'), 1


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
